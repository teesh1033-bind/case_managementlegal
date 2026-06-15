#!/usr/bin/env python3
"""Generate UDM Industrial Attachment Report (DOCX)."""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt
from docx.oxml.ns import qn


# --- Student and host details ---
STUDENT_FULL_NAME = "Cocotte Chrinsley James"
STUDENT_SURNAME = "CHRINSLEY"
STUDENT_FIRST = "Cocotte"
STUDENT_ID = "I24006"
PROGRAMME = "BSc (Hons) Software Engineering (Full-Time)"
IA_PERIOD = "27 April 2026 – 18 July 2026 (12 weeks)"
HOST_ORGANISATION = "Active Service"
HOST_TYPE = "IT and software services organisation"
DEPARTMENT = "Software Development / Information Technology"
MENTOR_NAME = "Mrs Mohaboob"
MENTOR_TITLE = "Industrial Attachment Mentor"
MENTOR_EMAIL = "As provided by Active Service"
MENTOR_PHONE = "As provided by Active Service"
LOCATION = "Mauritius"


def set_document_defaults(doc: Document) -> None:
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.space_after = Pt(6)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Times New Roman"
        run.font.color.rgb = None


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.bold = bold
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE


def build_report() -> Document:
    doc = Document()
    set_document_defaults(doc)

    # Cover page
    for _ in range(6):
        doc.add_paragraph()
    add_para(doc, "Université des Mascareignes", bold=True)
    add_para(doc, "Faculty of Information and Communication Technology")
    add_para(doc, "Department of Software Engineering")
    doc.add_paragraph()
    add_para(doc, "INDUSTRIAL ATTACHMENT REPORT", bold=True)
    doc.add_paragraph()
    add_para(doc, f"Student: {STUDENT_FULL_NAME}")
    add_para(doc, f"Student ID: {STUDENT_ID}")
    add_para(doc, f"Programme: {PROGRAMME}")
    add_para(doc, f"Host Organisation: {HOST_ORGANISATION}")
    add_para(doc, f"Attachment Period: {IA_PERIOD}")
    add_para(doc, f"Date of Submission: {date.today().strftime('%d %B %Y')}")
    doc.add_page_break()

    # Abstract
    add_heading(doc, "Abstract", 1)
    add_para(
        doc,
        "This report documents the industrial attachment undertaken as part of the BSc (Hons) "
        "Software Engineering programme at Université des Mascareignes. The attachment was completed "
        f"at {HOST_ORGANISATION} from {IA_PERIOD}. The primary focus was participation in the "
        "analysis, design, implementation, testing, and maintenance of web-based software, "
        "including a legal case management platform (LegalPro) used to support lawyers, "
        "administrators, and clients. Work included full-stack PHP development, MySQL database "
        "integration, responsive user-interface design, bug fixing, merge conflict resolution, "
        "and collaboration within a small development team using Git version control. "
        "The report describes the host organisation, tasks performed, technical challenges, "
        "and personal learning outcomes while respecting confidentiality requirements."
    )
    doc.add_page_break()

    # Table of contents placeholder
    add_heading(doc, "Table of Contents", 1)
    toc = [
        "1. Introduction",
        "2. Industrial Attachment Host",
        "3. Industrial Attachment Site",
        "4. Work Performed",
        "5. Experience, Appreciation, and Personal Remarks",
        "6. Conclusion",
        "7. References",
        "8. Appendix",
    ]
    for line in toc:
        add_para(doc, line)
    add_para(doc, "(In Word: click References → Table of Contents → Automatic Table 1)")
    doc.add_page_break()

    # 1 Introduction
    add_heading(doc, "1. Introduction", 1)
    add_heading(doc, "1.1 Introduction", 2)
    add_para(
        doc,
        "Industrial Attachment (IA) is a mandatory component of the Software Engineering degree "
        "that bridges academic learning and professional practice. During this attachment I was "
        "integrated into a software development environment where I contributed to real projects "
        "rather than isolated classroom exercises. The experience exposed me to software life-cycle "
        "activities including requirements understanding, implementation, debugging, code review "
        "through collaboration, and iterative UI improvement based on user feedback."
    )
    add_heading(doc, "1.2 Aim and Objectives", 2)
    add_para(doc, "The aims of this attachment were to:")
    add_bullets(
        doc,
        [
            "Apply software engineering principles to production-oriented web applications.",
            "Gain practical experience in PHP, MySQL, JavaScript, HTML/CSS, and Git workflows.",
            "Contribute to modules used by administrators, lawyers, and clients in a case management system.",
            "Develop professional skills: teamwork, communication, time management, and problem solving.",
            "Document work performed in a structured technical report aligned with UDM guidelines.",
        ],
    )
    add_heading(doc, "1.3 Brief Introduction to the Host Organisation", 2)
    add_para(
        doc,
        f"{HOST_ORGANISATION} operates in the field of information technology and software "
        "development. During my attachment at Active Service I worked on LegalPro, a legal-sector "
        "management web application that centralises case records, appointments, court dates, "
        "documents, billing, and client communication. Public-facing descriptions of the "
        "organisation’s activities are limited in this report; internal commercial data is omitted "
        "in accordance with confidentiality requirements."
    )
    add_heading(doc, "1.4 IA Mentor Contact Details", 2)
    add_para(doc, f"Name: {MENTOR_NAME}")
    add_para(doc, f"Position: {MENTOR_TITLE}")
    add_para(doc, f"Email: {MENTOR_EMAIL}")
    add_para(doc, f"Telephone: {MENTOR_PHONE}")

    # 5 Host (guidelines use section 5)
    add_heading(doc, "2. Industrial Attachment Host", 1)
    add_heading(doc, "2.1 Type and Group Structure", 2)
    add_para(
        doc,
        f"The host organisation is Active Service, described as a {HOST_TYPE}. "
        "Active Service provides technology solutions and supports software development "
        "projects for business clients in Mauritius."
    )
    add_heading(doc, "2.2 Nature of Activities", 2)
    add_para(
        doc,
        "Core activities include custom software development, web application maintenance, "
        "database design, and support for business users. The LegalPro platform supports "
        "case lifecycle management, scheduling, financial tracking, document storage, and "
        "role-based portals for admin staff, lawyers, and clients."
    )
    add_heading(doc, "2.3 Organisational Structure", 2)
    add_para(
        doc,
        "Reporting structure: Intern (Software Engineering) reports to Mrs Mohaboob (Mentor), "
        "who coordinates with the software development team and management at Active Service."
    )
    add_heading(doc, "2.4 Number of Employees and Clients", 2)
    add_para(
        doc,
        "Active Service operates as a small to medium-sized IT organisation. Exact employee "
        "and client figures are not disclosed in this report in accordance with confidentiality "
        "requirements of the host organisation."
    )
    add_heading(doc, "2.5 Turnover", 2)
    add_para(doc, "Financial turnover of the organisation is not disclosed in this report (N/A).")

    # 6 Site
    add_heading(doc, "3. Industrial Attachment Site", 1)
    add_heading(doc, "3.1 Location and General Characteristics", 2)
    add_para(doc, f"The attachment was carried out at Active Service, {LOCATION}. Work was performed in the software development department with standard developer workstations, local web server environment (Laragon), and access to Git for version control.")
    add_heading(doc, "3.2 Nature, Type and Amount of Work", 2)
    add_para(
        doc,
        "Work was organised as a continuous software project over the attachment period, "
        "with tasks assigned through team coordination and version control (Git/GitHub). "
        "Deliverables were incremental features and bug fixes merged into the main codebase."
    )
    add_heading(doc, "3.3 Work Team", 2)
    add_para(
        doc,
        "I worked within a small development team alongside other interns and developers "
        "(including collaborators identified in the Git repository). Tasks were divided by "
        "module (admin portal, lawyer portal, client portal) with regular merges and conflict resolution."
    )
    add_heading(doc, "3.4 Functional and Technical Organisation", 2)
    add_para(
        doc,
        "Technically the project uses a LAMP-style stack: PHP server pages, MySQL via PDO, "
        "Bootstrap/Argon Dashboard CSS, JavaScript for interactive components, and Laragon "
        "for local development on Windows. The codebase is structured into pages, shared includes "
        "(navigation, database, themes), and reusable libraries for domain logic."
    )

    # 7 Work performed - MAIN SECTION
    add_heading(doc, "4. Work Performed", 1)
    add_heading(doc, "4.1 Nature of the Work", 2)
    add_para(
        doc,
        "My work spanned analysis, implementation, testing, and maintenance of the LegalPro "
        "case management system. Although the product is legal-sector focused, the engineering "
        "skills are general: MVC-style page organisation, SQL queries, form validation, session "
        "management, and front-end usability. I also performed general IT-related activities "
        "such as environment setup, debugging production-like issues, and documentation through "
        "commit messages and code comments."
    )

    add_heading(doc, "4.2 Tasks, Execution and Technical Follow-up", 2)
    add_para(doc, "Major tasks and contributions during the attachment included:")
    add_bullets(
        doc,
        [
            "Admin portal: global UI improvements, sidebar navigation fixes (including layout at 100% zoom), dashboard enhancements, and integration of a configurable portal colour theme (CSS variables, dark mode).",
            "Lawyer portal: appointment scheduling features, lawyer availability management (recurring and date-specific slots), and lawyer-facing case/client views.",
            "Client portal: dashboard, case view, appointments booking with lawyer availability, court tracking with search, documents centre, and payments/invoices view.",
            "Appointment availability engine: implemented and debugged PHP logic in appointment_availability.php — including precedence of date-specific slots over recurring rules, validation of bookable time windows, and server-side booking validation.",
            "Client appointment booking UI: Flatpickr date picker with unavailable dates disabled, calendar positioning and styling aligned with portal theme, and time-slot dropdown showing available times.",
            "Finance module: quotation system (admin case view and client portal), invoice and payment pages, themed printable invoice/receipt/quotation documents using organisation colour presets.",
            "Case detail page: tabbed interface improvements for invoices, quotations, payments, and related case information.",
            "Tasks module: resolved Git merge conflicts and restored functional task listing/management.",
            "Court tracking: added search functionality for clients and lawyers.",
            "Documents module: improved filter/search form layout and accessibility (label alignment, touch-friendly controls).",
            "Removed deprecated ICS calendar export from client portal per product direction.",
            "Collaborative development: frequent Git merges, conflict resolution, and pair-style coordination with team members on shared modules.",
        ],
    )

    add_heading(doc, "4.3 Particular Problems", 2)
    add_bullets(
        doc,
        [
            "Merge conflicts in shared files (e.g. tasks.php, appointment_availability.php) caused parse errors; resolved by careful manual merging and syntax validation.",
            "Lawyer availability logic incorrectly blocked entire days when partial availability existed; fixed by giving date-specific slots precedence over recurring weekday rules.",
            "Calendar date picker misalignment in sticky layouts; solved with fixed positioning relative to the input field.",
            "Theme inconsistency across admin finance pages; addressed by centralising portal theme injection and shared finance CSS.",
            "Balancing feature requests with regression testing across three portals (admin, lawyer, client).",
        ],
    )

    add_heading(doc, "4.4 Work Conditions", 2)
    add_para(
        doc,
        "Work was performed on-site and/or remotely using a standard developer workstation, "
        "IDE (e.g. Cursor/VS Code), local web server (Laragon), and Git for source control. "
        "Tasks were assigned flexibly with milestone-style delivery through commits and merges."
    )

    add_heading(doc, "4.5 Nature of Relationships", 2)
    add_para(
        doc,
        "I maintained regular communication with my mentor and teammates for task clarification, "
        "code reviews informal through merge discussions, and user-feedback-driven UI adjustments. "
        "Professional conduct and respect for confidentiality were observed when handling case, "
        "client, and financial data in development databases."
    )

    add_heading(doc, "4.6 Analysis and Personal Remarks", 2)
    add_para(
        doc,
        "The attachment confirmed that real-world software engineering involves far more than "
        "writing new features: maintenance, debugging, merge management, and usability polish "
        "consume significant effort. Working on a multi-role system taught me how authentication, "
        "authorisation, and portal-specific UX must be designed consistently. The availability "
        "booking module in particular demonstrated how business rules (lawyer schedules, "
        "exceptions, holidays) must be modelled clearly in both database and application layers."
    )

    # 8 Experience
    add_heading(doc, "5. Experience, Appreciation, and Personal Remarks", 1)
    add_heading(doc, "5.1 Working Conditions", 2)
    add_para(
        doc,
        "Working hours followed the organisation’s normal business schedule, with flexibility "
        "for development tasks. Equipment included a developer workstation, internet access, "
        "local development tools (Laragon, PHP, MySQL), and Git for source control."
    )
    add_heading(doc, "5.2 Ease of Integration into the Team", 2)
    add_para(
        doc,
        "Integration was facilitated by using the same tools as permanent staff (Git, PHP, shared "
        "coding conventions) and by starting with smaller bug fixes before larger features."
    )
    add_heading(doc, "5.3 Relations with Co-workers, Management, and Mentor", 2)
    add_para(
        doc,
        f"My mentor, {MENTOR_NAME}, supervised my attachment at Active Service and provided "
        "guidance on task prioritisation, code quality, and professional conduct. Team members "
        "collaborated on shared branches and helped resolve merge conflicts during development."
    )
    add_heading(doc, "5.4 Critical Analysis", 2)
    add_para(
        doc,
        "The attachment was technically demanding and aligned well with Software Engineering "
        "learning outcomes: requirements from informal feedback, iterative delivery, and "
        "quality through testing. A structured ticketing system could further improve traceability."
    )
    add_heading(doc, "5.5 Difficulties Encountered", 2)
    add_bullets(
        doc,
        [
            "Understanding a large existing codebase without complete written specifications.",
            "Debugging date/time and timezone edge cases in appointment availability.",
            "Coordinating parallel changes from multiple developers on the same repository.",
            "Meeting UI expectations while preserving backward compatibility.",
        ],
    )

    # 9 Conclusion
    add_heading(doc, "6. Conclusion", 1)
    add_heading(doc, "6.1 Experience", 2)
    add_para(
        doc,
        "This industrial attachment provided hands-on experience in full-stack web development "
        "and professional software teamwork. I contributed to a production-oriented case "
        "management platform across admin, lawyer, and client interfaces, and strengthened "
        "skills in PHP, SQL, JavaScript, CSS, Git, and technical documentation."
    )
    add_heading(doc, "6.2 Critical Appraisal", 2)
    add_para(
        doc,
        "The experience met the IA objectives of the Software Engineering programme. "
        "I would benefit from earlier exposure to automated testing and formal code review "
        "checklists in future projects."
    )
    add_heading(doc, "6.3 Future Work", 2)
    add_bullets(
        doc,
        [
            "Automated unit and integration tests for appointment availability and billing modules.",
            "API layer for mobile client access to cases and appointments.",
            "Enhanced reporting and analytics on the financial summary dashboard.",
            "Continued accessibility and mobile responsiveness audits across all portals.",
        ],
    )

    # References
    add_heading(doc, "7. References", 1)
    add_bullets(
        doc,
        [
            "Université des Mascareignes (2026). Industrial Attachment Guidelines — Faculty of ICT, Department of Software Engineering, June 2026.",
            "PHP Group. PHP Manual. https://www.php.net/manual/en/",
            "Oracle Corporation. MySQL 8.0 Reference Manual. https://dev.mysql.com/doc/",
            "Bootstrap Team. Bootstrap Documentation. https://getbootstrap.com/docs/",
            "Flatpickr. Flatpickr Documentation. https://flatpickr.js.org/",
        ],
    )

    # Appendix
    add_heading(doc, "8. Appendix", 1)
    add_heading(doc, "Appendix A — Sample Module List (LegalPro)", 2)
    add_bullets(
        doc,
        [
            "Admin: Dashboard, Clients, Cases, Payments, Invoices, Finance Summary, Appointments, Court Tracking, Lawyers, Documents, AI Assistant.",
            "Lawyer: Dashboard, Cases, Clients, Availability, Appointments, Court Tracking, Tasks.",
            "Client: Dashboard, Cases, Appointments, Documents, Payments, Court Tracking, Profile, Settings.",
        ],
    )
    add_heading(doc, "Appendix B — Technologies Used", 2)
    add_para(doc, "PHP 8.x, MySQL, HTML5, CSS3, JavaScript, Bootstrap/Argon Dashboard, Git/GitHub, Laragon, PDO, Flatpickr.")
    add_heading(doc, "Appendix C — Submission Checklist", 2)
    add_bullets(
        doc,
        [
            "Generate final Table of Contents in Microsoft Word.",
            "Add organigram of Active Service if approved by mentor.",
            "Obtain signed Mentor Assessment sheet (Log Book page 21) with company seal.",
            f"Final PDF filename: {STUDENT_SURNAME}_{STUDENT_FIRST}_IA Report BSc SE FT 2026.pdf",
        ],
    )

    return doc


def main() -> None:
    out_dir = Path.home() / "Downloads"
    out_dir.mkdir(parents=True, exist_ok=True)
    base = f"{STUDENT_SURNAME}_{STUDENT_FIRST}_IA Report BSc SE FT 2026"
    docx_path = out_dir / f"{base}.docx"
    doc = build_report()
    doc.save(docx_path)
    print(f"Created: {docx_path}")
    print(f"PDF: {docx_path.with_suffix('.pdf')}")


if __name__ == "__main__":
    main()
